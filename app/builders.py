"""Построение отчётов: агрегаты считает БД, Excel собирается из готовых сумм (а не из 30 000 сырых строк)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import extract, func, select
from sqlalchemy.orm import Session

from app.models import Sale

MONTHS = ["Янв", "Фев", "Мар", "Апр", "Май", "Июн", "Июл", "Авг", "Сен", "Окт", "Ноя", "Дек"]
HEADER_FILL = PatternFill("solid", fgColor="1F2A5A")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def monthly_matrix(db: Session, year: int) -> tuple[list[str], list[list[int]], int]:
    """Выручка по месяцам и регионам: (регионы, 12 строк по регионам, число обработанных строк)."""
    rows = db.execute(
        select(extract("month", Sale.sold_on), Sale.region, func.sum(Sale.amount_cents), func.count())
        .where(extract("year", Sale.sold_on) == year)
        .group_by(extract("month", Sale.sold_on), Sale.region)
    ).all()
    regions = sorted({r[1] for r in rows})
    matrix = [[0] * len(regions) for _ in range(12)]
    processed = 0
    for month, region, total, cnt in rows:
        matrix[int(month) - 1][regions.index(region)] = int(total)
        processed += cnt
    return regions, matrix, processed


def top_products(db: Session, year: int, limit: int = 20) -> tuple[list[tuple[str, int, int]], int]:
    revenue = func.sum(Sale.amount_cents)
    rows = db.execute(
        select(Sale.product, revenue, func.count()).where(extract("year", Sale.sold_on) == year).group_by(Sale.product).order_by(revenue.desc(), Sale.product).limit(limit)
    ).all()
    processed = db.scalar(select(func.count()).where(extract("year", Sale.sold_on) == year)) or 0
    return [(p, int(r), int(c)) for p, r, c in rows], processed


def _style_header(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, Alignment(horizontal="center")


def build_sales_by_month(db: Session, year: int, path: Path, attempt: int = 1) -> int:
    regions, matrix, processed = monthly_matrix(db, year)
    wb = Workbook()
    ws = wb.active
    ws.title = f"Выручка {year}"
    ws.append(["Месяц", *regions, "Итого"])
    _style_header(ws, 1, len(regions) + 2)
    for i, row in enumerate(matrix, start=2):
        ws.append([MONTHS[i - 2], *[v / 100 for v in row], f"=SUM(B{i}:{chr(65 + len(regions))}{i})"])
    last = len(matrix) + 1
    ws.append(["Всего", *[f"=SUM({chr(66 + j)}2:{chr(66 + j)}{last})" for j in range(len(regions) + 1)]])
    for c in range(1, len(regions) + 3):
        ws.cell(row=last + 1, column=c).font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=len(regions) + 2):
        for cell in row:
            cell.number_format = '#,##0.00'
    ws.column_dimensions["A"].width = 12
    for j in range(len(regions) + 1):
        ws.column_dimensions[chr(66 + j)].width = 16
    chart = BarChart()
    chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
    chart.title, chart.height, chart.width = f"Выручка по месяцам и регионам, {year}", 9, 22
    chart.add_data(Reference(ws, min_col=2, max_col=len(regions) + 1, min_row=1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    ws.add_chart(chart, f"A{last + 4}")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return processed


def build_top_products(db: Session, year: int, path: Path, attempt: int = 1) -> int:
    rows, processed = top_products(db, year)
    wb = Workbook()
    ws = wb.active
    ws.title = f"Топ товаров {year}"
    ws.append(["#", "Товар", "Выручка", "Продаж"])
    _style_header(ws, 1, 4)
    for i, (name, revenue, count) in enumerate(rows, start=1):
        ws.append([i, name, revenue / 100, count])
        ws.cell(row=i + 1, column=3).number_format = '#,##0.00'
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 16
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return processed


def build_flaky(db: Session, year: int, path: Path, attempt: int = 1) -> int:
    """Демонстрация повторных попыток: источник данных «не отвечает» первые две попытки."""
    if attempt <= 2:
        raise ConnectionError("data warehouse is temporarily unavailable")
    return build_top_products(db, year, path, attempt)


BUILDERS = {"sales_by_month": build_sales_by_month, "top_products": build_top_products, "flaky_demo": build_flaky}
