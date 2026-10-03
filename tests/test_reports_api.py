import io
from pathlib import Path

from openpyxl import load_workbook

from app.db import SessionLocal
from app.models import Report
from tests.conftest import ROWS, make_report


def order(client, kind="sales_by_month", year=2025):
    r = client.post("/reports", json={"kind": kind, "year": year})
    assert r.status_code in (200, 202), r.text
    return r.json()


def test_report_is_built_and_reported_as_done(client):
    r = order(client)
    done = client.get(f"/reports/{r['id']}").json()
    assert done["status"] == "done" and done["progress"] == 100 and done["rows_processed"] == ROWS
    assert done["error"] is None and done["duration_s"] is not None and done["attempts"] == 1


def test_sales_by_month_excel_content(client, total_cents):
    rid = order(client)["id"]
    r = client.get(f"/reports/{rid}/download")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.title == "Выручка 2025"
    assert [c.value for c in ws[1]] == ["Месяц", "East", "North", "South", "West", "Итого"]
    # 12 месяцев + строка итогов; сумма значений по регионам равна сумме продаж в БД
    assert ws.max_row == 14 and ws["A14"].value == "Всего"
    cells = [ws.cell(row=r_, column=c).value for r_ in range(2, 14) for c in range(2, 6)]
    assert round(sum(cells) * 100) == total_cents
    assert str(ws["F2"].value).startswith("=SUM(") and str(ws["B14"].value) == "=SUM(B2:B13)"  # итоги - живые формулы Excel
    assert len(ws._charts) == 1                                                                 # диаграмма на листе


def test_top_products_sorted_and_correct(client):
    rid = order(client, "top_products")["id"]
    ws = load_workbook(io.BytesIO(client.get(f"/reports/{rid}/download").content)).active
    revenue = [ws.cell(row=i, column=3).value for i in range(2, ws.max_row + 1)]
    assert revenue == sorted(revenue, reverse=True) and ws["A2"].value == 1 and ws.max_row == 13  # 12 товаров + шапка
    assert sum(ws.cell(row=i, column=4).value for i in range(2, ws.max_row + 1)) == ROWS


def test_year_without_data_gives_empty_but_valid_report(client):
    r = order(client, year=2010)
    done = client.get(f"/reports/{r['id']}").json()
    assert done["status"] == "done" and done["rows_processed"] == 0


def test_validation(client):
    assert client.post("/reports", json={"kind": "nope"}).status_code == 422
    assert client.post("/reports", json={"kind": "top_products", "year": 1800}).status_code == 422


def test_not_ready_and_unknown(client):
    rid = make_report(status="running")
    assert client.get(f"/reports/{rid}/download").status_code == 409
    assert client.get(f"/reports/{rid}/preview").status_code == 409
    assert client.get("/reports/9999").status_code == 404
    assert client.get("/reports/9999/download").status_code == 404


def test_duplicate_order_returns_the_active_report(client):
    rid = make_report(kind="top_products", year=2025, status="running")
    r = client.post("/reports", json={"kind": "top_products", "year": 2025})
    assert r.status_code == 200 and r.json()["id"] == rid
    with SessionLocal() as db:
        assert db.query(Report).count() == 1


def test_finished_report_does_not_block_a_new_order(client):
    first = order(client, "top_products")["id"]
    second = order(client, "top_products")["id"]
    assert second != first


def test_delete_removes_row_and_file(client):
    rid = order(client, "top_products")["id"]
    path = Path(client.get(f"/reports/{rid}").json() and _path(rid))
    assert path.exists()
    assert client.delete(f"/reports/{rid}").status_code == 204
    assert not path.exists() and client.get(f"/reports/{rid}").status_code == 404


def _path(rid: int) -> str:
    with SessionLocal() as db:
        return db.get(Report, rid).file_path


def test_cancel_marks_active_report(client):
    rid = make_report(status="running")
    assert client.delete(f"/reports/{rid}").status_code == 204
    assert client.get(f"/reports/{rid}").json()["status"] == "cancelled"


def test_list_newest_first(client):
    a, b = order(client, "top_products")["id"], order(client, "sales_by_month")["id"]
    assert [r["id"] for r in client.get("/reports").json()] == [b, a]


def test_preview_page_shows_totals(client, total_cents):
    rid = order(client)["id"]
    html = client.get(f"/reports/{rid}/preview").text
    assert "Выручка по месяцам и регионам, 2025" in html and f"{total_cents / 100:,.2f}".replace(",", " ") in html
    top = order(client, "top_products")["id"]
    assert "Product" in client.get(f"/reports/{top}/preview").text
