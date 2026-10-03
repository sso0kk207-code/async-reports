"""Демо: python scripts/demo.py http://127.0.0.1:8000   (сервис запущен через scripts/dev.py: API + Celery-воркер)"""
import io
import time

from openpyxl import load_workbook

from _demo import client, section, show

section("1. Заказываем три отчёта. API отвечает мгновенно (202), работа уходит воркеру")
ids = {}
for kind in ("sales_by_month", "top_products", "flaky_demo"):
    t = time.perf_counter()
    r = show("POST", "/reports", json_body={"kind": kind, "year": 2025}, max_lines=0)
    ids[kind] = r.json()["id"]
    print(f"   ответ за {(time.perf_counter() - t) * 1000:.0f} мс, id={r.json()['id']}, status={r.json()['status']}")
    print()

section("2. Повторный заказ того же отчёта не плодит дубликат тяжёлой работы")
r = show("POST", "/reports", json_body={"kind": "sales_by_month", "year": 2025}, max_lines=0)
print(f"   вернулся тот же отчёт: id={r.json()['id']} (200 вместо 202)")
print()

section("3. Следим за прогрессом (flaky_demo два раза падает и повторяется сам)")
started, last = time.perf_counter(), {}
while time.perf_counter() - started < 60:
    states = {k: client.get(f"/reports/{i}").json() for k, i in ids.items()}
    line = {k: f"{s['status']} {s['progress']}% (попытка {s['attempts']})" for k, s in states.items()}
    if line != last:
        print(f"t={time.perf_counter() - started:4.1f}с  " + "  |  ".join(f"{k}: {v}" for k, v in line.items()))
        last = line
    if all(s["status"] in ("done", "failed") for s in states.values()):
        break
    time.sleep(0.3)
print()

section("4. Итог по отчётам")
for k, i in ids.items():
    s = client.get(f"/reports/{i}").json()
    print(f"#{i} {k:<15} {s['status']:<5} строк обработано: {s['rows_processed']:>6}   построен за {s['duration_s']} с   попыток: {s['attempts']}")
print()

section("5. Скачиваем Excel и заглядываем внутрь")
r = client.get(f"/reports/{ids['sales_by_month']}/download")
print(f"$ curl -O /reports/{ids['sales_by_month']}/download   ->  HTTP {r.status_code}, {len(r.content):,} байт, {r.headers['content-disposition']}")
ws = load_workbook(io.BytesIO(r.content)).active
print(f"лист «{ws.title}», строк: {ws.max_row}, диаграмм: {len(ws._charts)}")
for row in ws.iter_rows(min_row=1, max_row=4, values_only=True):
    print("   ", [round(v, 2) if isinstance(v, float) else v for v in row])
print("   ...")
print("   ", [c.value for c in ws[ws.max_row]], " <- итоги это формулы Excel")
